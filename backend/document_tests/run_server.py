"""Create isolated browser-test fixtures and start the real parser."""
from io import BytesIO
from pathlib import Path
import sys

import fitz
from docx import Document
import uvicorn

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
directory = Path(__file__).resolve().parents[2] / "frontend" / ".source-test-fixtures"
directory.mkdir(exist_ok=True)
with fitz.open() as pdf:
    for text in ["Page one: study aim.", "Page two: interview results."]:
        page = pdf.new_page(); page.insert_text((72, 72), text)
    pdf.save(directory / "study.pdf")
with fitz.open() as pdf:
    pdf.new_page(); pdf.save(directory / "scan.pdf")
doc = Document(); doc.add_paragraph("Research method: interviews.")
doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table findings."
doc.save(directory / "study.docx")
(directory / "notes.txt").write_text("Original line one.\n\nOriginal line three.\n<script>untrusted()</script>", encoding="utf-8")
(directory / "broken.pdf").write_bytes(b"%PDF-invalid")
if __name__ == "__main__":
    # Browser tests exercise real uploads/retrieval/validation, but never spend credits.
    # This replacement exists only in the test runner, not in the production API.
    import json
    from document_service.main import app
    from document_service import chat
    from document_service.retrieval import retrieve

    async def fixture_completion(config, client, messages):
        data = json.loads(messages[1]["content"])
        ids = [item["id"] for item in data["evidence"] if item["text"] == "Page two: interview results."]
        if data["history"]:
            answer = "In plain language, this passage only indicates that the page covers interview results; it does not give the results themselves."
        else:
            answer = "The source identifies interview results as a topic on page two."
        return json.dumps({"answer": answer, "evidence_ids": ids[:1]})

    async def fixture_retrieval(question):
        return retrieve(question)

    chat.complete = fixture_completion
    chat.ChatConfig.from_env = lambda: None
    chat.retrieve_sources = fixture_retrieval
    uvicorn.run(app, host="127.0.0.1", port=8012)
