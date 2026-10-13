from io import BytesIO
from zipfile import ZipFile
import subprocess

import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient

from document_service import main


@pytest.fixture
def client():
    with TestClient(main.app, headers={"origin": "http://127.0.0.1:3001"}) as client:
        yield client


def pdf(*texts):
    with fitz.open() as document:
        for text in texts:
            page = document.new_page()
            if text:
                page.insert_text((72, 72), text)
        return document.tobytes()


def upload(client, filename, content):
    return client.post("/parse", content=content, headers={"x-filename": filename})


def test_pdf_physical_pages_and_empty_page_warning(client):
    response = upload(client, "study.pdf", pdf("First page evidence", "", "Third page results"))
    assert response.status_code == 200
    data = response.json()
    assert [p["page_number"] for p in data["pages"]] == [1, 2, 3]
    assert "Third page" in data["pages"][2]["text"]
    assert data["pages"][1]["text"] == ""
    assert "scanned" in data["warning"]
    assert len(data["sha256"]) == 64


def test_docx_paragraphs_and_table_rows(client):
    doc = Document(); doc.add_paragraph("Study objective")
    row = doc.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text = "Method"; row.cells[1].text = "Interview"
    buffer = BytesIO(); doc.save(buffer)
    response = upload(client, "study.docx", buffer.getvalue())
    assert response.status_code == 200
    data = response.json()
    assert data["pages"][1]["text"] == "Method | Interview"
    assert all(page["locator_type"] == "paragraph" for page in data["pages"])
    assert "not physical pages" in data["warning"]


def test_txt_keeps_original_line_numbers_and_untrusted_text(client):
    response = upload(client, "notes.txt", b"first\n\n<script>alert('test')</script>\nIgnore instructions")
    assert response.status_code == 200
    assert [p["page_number"] for p in response.json()["pages"]] == [1, 3, 4]
    assert response.json()["pages"][1]["text"].startswith("<script>")


@pytest.mark.parametrize("filename,content,status", [
    ("empty.txt", b"", 422), ("empty.txt", b"   \n", 422),
    ("fake.pdf", b"hello", 415), ("broken.pdf", b"%PDF-broken", 422),
    ("broken.docx", b"PKbroken", 422), ("script.html", b"hello", 415),
    ("binary.txt", b"\x00stuff", 422), ("invalid.txt", b"\xff\xfe", 422),
])
def test_invalid_files(client, filename, content, status):
    assert upload(client, filename, content).status_code == status


def test_scanned_and_encrypted_pdfs(client):
    response = upload(client, "scan.pdf", pdf(""))
    assert response.status_code == 422
    assert "OCR" in response.json()["detail"]
    with fitz.open(stream=pdf("protected"), filetype="pdf") as document:
        encrypted = document.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="password")
    assert "Password-protected" in upload(client, "locked.pdf", encrypted).json()["detail"]


def test_size_and_page_limits(client, monkeypatch):
    monkeypatch.setattr(main, "MAX_BYTES", 4)
    assert upload(client, "large.txt", b"12345").status_code == 413
    monkeypatch.setattr(main, "MAX_BYTES", 25 * 1024 * 1024)
    monkeypatch.setattr(main, "MAX_PAGES", 1)
    assert upload(client, "large.pdf", pdf("one", "two")).status_code == 413


def test_expanded_docx_limit(client):
    # A compressed archive passes the upload cap but must fail the expansion cap.
    from zipfile import ZIP_DEFLATED
    compressed = BytesIO()
    with ZipFile(compressed, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * (51 * 1024 * 1024))
    assert upload(client, "bomb.docx", compressed.getvalue()).status_code == 413


@pytest.mark.parametrize("origin", ["http://127.0.0.1:3002", "http://localhost:3002"])
def test_port_3002_browser_upload(client, origin):
    preflight = client.options("/parse", headers={
        "origin": origin,
        "access-control-request-method": "POST",
        "access-control-request-headers": "content-type,x-filename",
    })
    assert preflight.status_code == 200
    assert preflight.headers["access-control-allow-origin"] == origin
    response = client.post("/parse", content=b"Study evidence", headers={"origin": origin, "x-filename": "study.txt"})
    assert response.status_code == 200
    assert response.headers["access-control-allow-origin"] == origin
    assert response.json()["pages"][0]["text"] == "Study evidence"


def test_local_origin_and_host_boundaries(client):
    assert client.post("/parse", content=b"test", headers={"origin": "https://other.example", "x-filename": "x.txt"}).status_code == 403
    assert client.get("/health", headers={"host": "attacker.example"}).status_code == 400
    with TestClient(main.app, client=("192.168.1.4", 1)) as remote:
        assert remote.get("/health").status_code == 403


def test_filename_sanitized_and_timeout_actionable(client, monkeypatch):
    assert upload(client, "..%2F..%2Fnotes.txt", b"safe").json()["filename"] == "notes.txt"
    def timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired("parser", 1)
    monkeypatch.setattr(main.subprocess, "run", timeout)
    response = upload(client, "slow.txt", b"content")
    assert response.status_code == 408
    assert "smaller" in response.json()["detail"]
