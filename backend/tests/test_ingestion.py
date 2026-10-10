"""Tests for page extraction and page-safe, overlapping token chunks."""

from io import BytesIO

import fitz
import pytest
from docx import Document
from fastapi import HTTPException

from app.services.chunking import chunk_pages
from app.services.parsing import ExtractedPage, extract_pages


def _pdf(*pages: str) -> bytes:
    pdf = fitz.open()
    for text in pages:
        page = pdf.new_page()
        page.insert_text((72, 72), text)
    return pdf.tobytes()


def test_pdf_extraction_preserves_one_based_page_numbers() -> None:
    pages, warning = extract_pages("sample.pdf", _pdf("Aim: examine sleep", "Result: sleep improved"))
    assert [page.page_number for page in pages] == [1, 2]
    assert "Aim" in pages[0].text
    assert warning is None


def test_image_only_pdf_is_rejected_with_actionable_ocr_warning() -> None:
    pdf = fitz.open()
    pdf.new_page()
    with pytest.raises(HTTPException) as exc:
        extract_pages("scanned.pdf", pdf.tobytes())
    assert "OCR is not supported" in exc.value.detail


def test_zero_byte_source_is_rejected() -> None:
    with pytest.raises(HTTPException) as exc:
        extract_pages("blank.txt", b"")
    assert exc.value.status_code == 422


def test_txt_extraction_uses_documented_logical_page() -> None:
    pages, warning = extract_pages("notes.txt", b"Evidence is described here.")
    assert pages == [pages[0]]
    assert pages[0].page_number == 1
    assert pages[0].locator_type == "line"
    assert "line numbers" in warning


def test_docx_extraction_warns_that_page_one_is_logical() -> None:
    document = Document()
    document.add_paragraph("This is a formatted thesis note.")
    output = BytesIO()
    document.save(output)
    pages, warning = extract_pages("notes.docx", output.getvalue())
    assert pages[0].page_number == 1
    assert pages[0].text == "This is a formatted thesis note."
    assert pages[0].locator_type == "paragraph"
    assert "paragraph numbers" in warning


def test_chunker_enforces_token_limit_overlap_and_source_page() -> None:
    text = " ".join(f"academicword{i}" for i in range(160))
    chunks = chunk_pages([ExtractedPage(page_number=7, text=text)], "paper.pdf", chunk_size_tokens=30, overlap_tokens=5)
    assert len(chunks) > 1
    assert all(chunk["page"] == 7 and chunk["filename"] == "paper.pdf" for chunk in chunks)
    assert all(len(chunk["text"]) < 200 for chunk in chunks)


def test_chunker_rejects_invalid_overlap() -> None:
    with pytest.raises(ValueError):
        chunk_pages([], "paper.pdf", chunk_size_tokens=4, overlap_tokens=4)


def test_docx_table_and_corrupt_zip():
    document = Document()
    document.add_paragraph("Aim")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Participants"
    table.cell(0, 1).text = "42"
    buffer = BytesIO()
    document.save(buffer)
    pages, _ = extract_pages("table.docx", buffer.getvalue())
    assert pages[1].text == "Participants | 42"
    assert pages[1].page_number == 2
    with pytest.raises(HTTPException) as exc:
        extract_pages("bad.docx", b"PKcorrupt")
    assert exc.value.status_code == 422


def test_text_preserves_line_numbers_and_literal_special_tokens():
    pages, _ = extract_pages("lines.txt", b"first\n\n<|endoftext|> last")
    assert [page.page_number for page in pages] == [1, 3]
    chunks = chunk_pages(pages, "lines.txt")
    assert chunks[-1]["page"] == 3
    assert chunks[-1]["locator_type"] == "line"


def test_pdf_page_limit(monkeypatch):
    from app.core.config import settings
    monkeypatch.setattr(settings, "max_pdf_pages", 1)
    with pytest.raises(HTTPException) as exc:
        extract_pages("large.pdf", _pdf("one", "two"))
    assert exc.value.status_code == 413
